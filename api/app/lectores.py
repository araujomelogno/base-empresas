"""Lectura de bases en .xlsx, .csv y .sav a una lista de filas (dict columna → valor)."""
import csv
import io
import os
import tempfile

from fastapi import HTTPException

from .util import texto


def leer_tabla(datos: bytes, tipo: str) -> tuple[list[str], list[dict], dict]:
    """Devuelve (columnas, filas, metadatos). Los metadatos solo vienen para .sav."""
    if tipo == "csv":
        return (*_csv(datos), {})
    if tipo == "xlsx":
        return (*_xlsx(datos), {})
    if tipo == "sav":
        return leer_sav(datos)
    raise HTTPException(415, f"Tipo no soportado: {tipo}")


def _decodificar(datos: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return datos.decode(enc)
        except UnicodeDecodeError:
            continue
    raise HTTPException(422, "No se pudo leer la codificación del CSV")


def _csv(datos: bytes):
    txt = _decodificar(datos)
    muestra = txt[:20000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t|")
    except csv.Error:
        dialecto = csv.excel
        dialecto.delimiter = ";" if muestra.count(";") > muestra.count(",") else ","
    lector = csv.reader(io.StringIO(txt), dialecto)
    filas = list(lector)
    if not filas:
        raise HTTPException(422, "El archivo está vacío")
    columnas = _columnas(filas[0])
    return columnas, [dict(zip(columnas, _rellenar(f, len(columnas)))) for f in filas[1:] if any(c.strip() for c in f)]


def _xlsx(datos: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(422, f"No se pudo abrir el Excel: {e}") from e
    ws = wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    try:
        encabezado = next(filas)
    except StopIteration:
        raise HTTPException(422, "El Excel está vacío") from None
    columnas = _columnas(["" if c is None else str(c) for c in encabezado])
    salida = []
    for f in filas:
        if f is None or all(c is None or str(c).strip() == "" for c in f):
            continue
        salida.append(dict(zip(columnas, _rellenar(list(f), len(columnas)))))
    wb.close()
    return columnas, salida


def leer_sav(datos: bytes) -> tuple[list[str], list[dict], dict]:
    """Lee un .sav con pyreadstat (códigos crudos, sin aplicar etiquetas)."""
    import pyreadstat

    fd, ruta = tempfile.mkstemp(suffix=".sav")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(datos)
        df, meta = pyreadstat.read_sav(ruta, apply_value_formats=False, user_missing=True)
    except Exception as e:
        raise HTTPException(422, f"No se pudo leer el .sav: {e}") from e
    finally:
        os.unlink(ruta)
    columnas = list(df.columns)
    filas = []
    for rec in df.to_dict(orient="records"):
        filas.append({k: _valor_sav(v) for k, v in rec.items()})
    etiquetas = dict(zip(meta.column_names, meta.column_labels or [None] * len(columnas)))
    valores = {}
    for var, nombre_set in (meta.variable_to_label or {}).items():
        valores[var] = {_clave(k): v for k, v in meta.value_labels.get(nombre_set, {}).items()}
    return columnas, filas, {"etiquetas": etiquetas, "valores": valores}


def _valor_sav(v):
    if v is None:
        return None
    if isinstance(v, float):
        if v != v:
            return None
        return int(v) if v.is_integer() else v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    s = str(v)
    return s if s.strip() else None


def _clave(k) -> str:
    """Clave de un código de valor como texto: 1.0 -> '1'."""
    if isinstance(k, float) and k.is_integer():
        k = int(k)
    return str(k)


def clave_valor(v) -> str | None:
    t = texto(v)
    return None if t is None else _clave(v if not isinstance(v, str) else t)


def _columnas(encabezado: list[str]) -> list[str]:
    vistos: dict[str, int] = {}
    cols = []
    for i, c in enumerate(encabezado):
        c = (c or "").strip() or f"columna_{i + 1}"
        if c in vistos:
            vistos[c] += 1
            c = f"{c}_{vistos[c]}"
        else:
            vistos[c] = 1
        cols.append(c)
    return cols


def _rellenar(fila: list, n: int) -> list:
    fila = list(fila)[:n]
    return fila + [None] * (n - len(fila))
