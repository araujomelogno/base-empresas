"""Utilidades comunes de los scripts de la fase 0.

Reusa los módulos puros del backend (api/app): telefonos, rut, lectores y sav. No hace
falta base de datos: importar app.sav carga app.db, pero el pool se abre recién al usarlo.
"""
import csv
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

FASE0 = Path(__file__).resolve().parent
RAIZ = FASE0.parent
SALIDA = FASE0 / "salida"
EJEMPLO = SALIDA / "ejemplo"
SEED_SQL = RAIZ / "db" / "empresas" / "0003_seed.sql"
INTERNOS_CSV = FASE0 / "datos" / "destinos_internos.csv"

if str(RAIZ / "api") not in sys.path:
    sys.path.insert(0, str(RAIZ / "api"))

from app import lectores  # noqa: E402
from app import rut as rutmod  # noqa: E402
from app import telefonos as telmod  # noqa: E402
from app.util import razon_normalizada, texto  # noqa: E402,F401


class ErrorFase0(Exception):
    """Error de insumos (archivo ilegible, columna inexistente…)."""


# ─── Lectura de archivos ───

def tipo_archivo(ruta) -> str:
    ext = Path(ruta).suffix.lower()
    if ext in (".csv", ".txt", ".tsv"):
        return "csv"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext == ".sav":
        return "sav"
    raise ErrorFase0(f"{ruta}: formato no soportado ({ext}); se aceptan .csv, .xlsx y .sav")


def leer(ruta, hoja: str | None = None) -> tuple[list[str], list[dict], dict]:
    """Lee un .csv/.xlsx/.sav con app.lectores. Devuelve (columnas, filas, meta)."""
    ruta = Path(ruta)
    if not ruta.exists():
        raise ErrorFase0(f"No existe el archivo {ruta}")
    tipo = tipo_archivo(ruta)
    try:
        if tipo == "xlsx" and hoja:
            return (*_leer_hoja(ruta, hoja), {})
        return lectores.leer_tabla(ruta.read_bytes(), tipo)
    except ErrorFase0:
        raise
    except Exception as e:  # HTTPException de lectores u otros
        raise ErrorFase0(f"{ruta.name}: {getattr(e, 'detail', e)}") from e


def _leer_hoja(ruta: Path, hoja: str):
    """Como lectores._xlsx pero eligiendo la hoja por nombre (sin distinguir mayúsculas)."""
    from openpyxl import load_workbook

    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        nombres = {n.strip().lower(): n for n in wb.sheetnames}
        nombre = nombres.get(hoja.strip().lower())
        if nombre is None:
            raise ErrorFase0(f"{ruta.name}: no tiene la hoja '{hoja}' (hojas: {', '.join(wb.sheetnames)})")
        filas = wb[nombre].iter_rows(values_only=True)
        encabezado = next(filas, None)
        if encabezado is None:
            raise ErrorFase0(f"{ruta.name}: la hoja '{nombre}' está vacía")
        columnas = lectores._columnas(["" if c is None else str(c) for c in encabezado])
        salida = [dict(zip(columnas, lectores._rellenar(list(f), len(columnas)))) for f in filas
                  if f is not None and any(c is not None and str(c).strip() for c in f)]
        return columnas, salida
    finally:
        wb.close()


# ─── Columnas ───

def norm(s) -> str:
    """Texto comparable: minúsculas, sin tildes, solo letras y números."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]", "", s)


def buscar_columna(columnas, candidatos, contiene: bool = False) -> str | None:
    """Primera columna cuyo nombre normalizado coincide con un candidato (o lo contiene)."""
    idx = {norm(c): c for c in columnas}
    for cand in candidatos:
        if norm(cand) in idx:
            return idx[norm(cand)]
    if contiene:
        for cand in candidatos:
            for n, c in idx.items():
                if norm(cand) and norm(cand) in n:
                    return c
    return None


def exigir_columna(columnas, nombre, archivo) -> str:
    col = buscar_columna(columnas, [nombre])
    if col is None:
        raise ErrorFase0(f"{archivo}: no tiene la columna '{nombre}' (columnas: {', '.join(columnas)})")
    return col


_RE_TEL = re.compile(r"tel|fono|cel|movil|phone|whatsapp|contacto_?num")


def detectar_columnas_tel(columnas, filas, muestra: int = 300) -> list[str]:
    """Columnas de teléfono: por nombre y, si no hay, por contenido (≥ 50% normalizable)."""
    por_nombre = [c for c in columnas if _RE_TEL.search(norm(c)) and "rut" not in norm(c)]
    if por_nombre:
        return por_nombre
    salida = []
    for c in columnas:
        vals = [f.get(c) for f in filas[:muestra] if texto(f.get(c))]
        if vals and sum(1 for v in vals if any(telmod.normalizar(p).discable for p in telmod.separar(v))) >= len(vals) / 2:
            if sum(1 for v in vals if rutmod.es_valido(v)) < len(vals) / 2:
                salida.append(c)
    return salida


def detectar_columna_rut(columnas, filas, muestra: int = 300) -> str | None:
    col = buscar_columna(columnas, ["rut", "EC_RUT", "rut_empresa", "rutempresa", "nro_rut", "ruc"])
    if col:
        return col
    for c in columnas:
        if "rut" in norm(c):
            return c
    for c in columnas:
        vals = [f.get(c) for f in filas[:muestra] if texto(f.get(c))]
        if vals and sum(1 for v in vals if rutmod.es_valido(v)) >= len(vals) / 2:
            return c
    return None


def detectar_columna_razon(columnas) -> str | None:
    return buscar_columna(columnas, ["razon_social", "razón social", "razon", "denominacion", "empresa", "nombre",
                                     "nombre_empresa"], contiene=True)


def clave_id(v) -> str | None:
    """id_item comparable: 1001, 1001.0 y '1001' dan '1001'."""
    t = texto(v)
    if t is None:
        return None
    try:
        f = float(t)
        return str(int(f)) if f.is_integer() else t
    except ValueError:
        return t


def telefonos_de(fila: dict, columnas_tel, prefijos=()) -> list:
    """Teléfonos normalizados de una fila (una celda puede traer varios)."""
    salida = []
    for c in columnas_tel:
        for parte in telmod.separar(fila.get(c)):
            salida.append(telmod.normalizar(parte, prefijos))
    return salida


# ─── Configuración por defecto (la misma del seed de la base) ───

def config_seed() -> dict:
    """Lee los parámetros de la versión 1 de la configuración desde db/empresas/0003_seed.sql."""
    defecto = {"destinos_internos": [r"^\d{1,5}$", r"^\*", r"^#"], "prefijos_salida": [], "umbral_intermediario": 5}
    try:
        sql = SEED_SQL.read_text(encoding="utf-8")
        m = re.search(r"INSERT INTO config_version.*?'(\{.*?\})'::jsonb", sql, re.S)
        return json.loads(m.group(1)) if m else defecto
    except (OSError, ValueError):
        return defecto


def destinos_internos(extra_csv=INTERNOS_CSV, extra=()) -> list[str]:
    """Patrones del seed + los de fase0/datos/destinos_internos.csv (filas con patrón) + extra."""
    patrones = list(config_seed().get("destinos_internos") or [])
    if extra_csv and Path(extra_csv).exists():
        with open(extra_csv, encoding="utf-8-sig", newline="") as f:
            for fila in csv.DictReader(f):
                p = (fila.get("patron") or "").strip()
                if p and p not in patrones:
                    patrones.append(p)
    for p in extra or ():
        if p and p not in patrones:
            patrones.append(p)
    return patrones


# ─── Salida ───

def pct(a, b) -> float:
    return round(100 * a / b, 1) if b else 0.0


def md_tabla(encabezados, filas) -> str:
    def celda(v):
        return str("" if v is None else v).replace("|", "\\|").replace("\n", " ")
    lineas = ["| " + " | ".join(encabezados) + " |", "| " + " | ".join("---" for _ in encabezados) + " |"]
    lineas += ["| " + " | ".join(celda(v) for v in f) + " |" for f in filas]
    return "\n".join(lineas)


def escribir_tabla(ruta, encabezados, filas) -> Path:
    """Escribe un detalle en .csv (separador ;, UTF-8 con BOM para Excel) o .xlsx."""
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    if ruta.suffix.lower() == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(list(encabezados))
        for f in filas:
            ws.append(list(f))
        wb.save(ruta)
    else:
        with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(encabezados)
            w.writerows(filas)
    return ruta


def escribir(ruta, contenido: str) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(contenido, encoding="utf-8")
    return ruta


def decision(go: bool | None) -> str:
    return {True: "GO", False: "NO GO", None: "sin datos"}[go]


def rel(ruta) -> str:
    """Ruta relativa a la raíz del repo, para los informes."""
    try:
        return os.path.relpath(Path(ruta).resolve(), RAIZ)
    except ValueError:
        return str(ruta)


# ─── Lista de teléfonos del estudio (insumo de A y B) ───

class Lista:
    """Lista del estudio indexada por teléfono E.164, RUT e id_item.

    La "empresa" de cada fila es su RUT válido; si no tiene, su id_item; si no, el número de fila.
    """

    def __init__(self, ruta, col_tel=None, col_rut=None, col_id=None, hoja=None, prefijos=()):
        self.ruta = Path(ruta)
        columnas, filas, _ = leer(ruta, hoja)
        self.filas = filas
        self.col_tel = [exigir_columna(columnas, c, self.ruta.name) for c in col_tel] if col_tel \
            else detectar_columnas_tel(columnas, filas)
        self.col_rut = exigir_columna(columnas, col_rut, self.ruta.name) if col_rut else detectar_columna_rut(columnas, filas)
        self.col_id = exigir_columna(columnas, col_id, self.ruta.name) if col_id \
            else buscar_columna(columnas, ["id_item", "EC_ID_ITEM", "iditem", "id"])
        if not self.col_tel:
            raise ErrorFase0(f"{self.ruta.name}: no se encontró columna de teléfono; indicala con --col-tel")
        self.por_tel: dict[str, set] = {}
        self.por_rut: dict[str, list[int]] = {}
        self.por_id: dict[str, int] = {}
        self.empresa: list[str] = []
        self.tels_fila: list[set] = []
        self.n_tel = 0
        self.tel_invalidos: dict[str, int] = {}
        self.ruts_invalidos = 0
        for i, f in enumerate(filas):
            r = rutmod.validar(f.get(self.col_rut))[0] if self.col_rut else None
            if self.col_rut and r is None and texto(f.get(self.col_rut)):
                self.ruts_invalidos += 1
            idv = clave_id(f.get(self.col_id)) if self.col_id else None
            clave = r or (f"id:{idv}" if idv else f"fila:{i + 2}")
            self.empresa.append(clave)
            if r:
                self.por_rut.setdefault(r, []).append(i)
            if idv:
                self.por_id[idv] = i
            tels = set()
            for t in telefonos_de(f, self.col_tel, prefijos):
                self.n_tel += 1
                if t.e164:
                    tels.add(t.e164)
                    self.por_tel.setdefault(t.e164, set()).add(clave)
                else:
                    self.tel_invalidos[t.tipo] = self.tel_invalidos.get(t.tipo, 0) + 1
            self.tels_fila.append(tels)

    @property
    def n_empresas(self) -> int:
        return len(set(self.empresa))

    def resumen(self) -> list[tuple]:
        return [
            ("Archivo", rel(self.ruta)),
            ("Filas", len(self.filas)),
            ("Empresas distintas", self.n_empresas),
            ("Columnas de teléfono", ", ".join(self.col_tel)),
            ("Columna RUT", self.col_rut or "—"),
            ("Columna id_item", self.col_id or "—"),
            ("Teléfonos leídos", self.n_tel),
            ("Teléfonos E.164 distintos", len(self.por_tel)),
            ("Teléfonos no normalizables", sum(self.tel_invalidos.values())),
        ]
